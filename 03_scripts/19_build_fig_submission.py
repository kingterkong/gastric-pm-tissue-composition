#!/usr/bin/env python3
"""Build the Functional & Integrative Genomics submission package."""
from pathlib import Path
import csv
import re
import shutil
import subprocess
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "06_manuscript" / "FIG_source"
OUT = ROOT / "06_manuscript" / "FIG_submission"
FIGOUT = OUT / "Figures"
PANDOC = Path("/Applications/RStudio.app/Contents/Resources/app/quarto/bin/tools/aarch64/pandoc")
SOFFICE = Path("/opt/homebrew/bin/soffice")
PDFTOPPM = Path("/opt/homebrew/bin/pdftoppm")
BIB = ROOT / "07_references" / "references.bib"
CSL = SRC / "springer-basic-author-date.csl"

OUT.mkdir(parents=True, exist_ok=True)
FIGOUT.mkdir(parents=True, exist_ok=True)

def read_csv(path):
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))

def esc(value):
    return str(value).replace("|", "\\|").replace("\n", " ")

def md_table(headers, rows, aligns=None):
    aligns = aligns or ["---"] * len(headers)
    out = ["| " + " | ".join(map(esc, headers)) + " |",
           "| " + " | ".join(aligns) + " |"]
    out += ["| " + " | ".join(esc(x) for x in row) + " |" for row in rows]
    return "\n".join(out)

def num(x, digits=3):
    if x in (None, "", "NA", "nan"):
        return "—"
    return f"{float(x):.{digits}f}"

def pval(x):
    if x in (None, "", "NA", "nan"):
        return "—"
    v = float(x)
    return "<0.001" if v < 0.001 else f"{v:.3f}"

# Main tables are deliberately compact enough for a portrait Word manuscript.
t1raw = read_csv(ROOT / "04_results" / "manuscript" / "Table1_sources.csv")
roles = {
    "GSE314812": "Primary paired bulk",
    "GSE237876": "External paired bulk / purity",
    "GSE183904": "Single-cell calibration",
    "GSE163558": "Single-cell calibration",
    "GSE308231": "Cell-level description",
    "GSE239676": "Fluid context",
    "GSE228598": "Metadata/context",
}
t1rows = []
for r in t1raw:
    other = []
    for key, label in [("Ascites", "ascites"), ("Lavage", "lavage"),
                       ("Other metastasis", "other metastasis"), ("Normal", "normal")]:
        if int(float(r[key])):
            other.append(f"{r[key]} {label}")
    t1rows.append([r["Source"], roles.get(r["Source"], "Context"), r["Specimens"],
                   r["Patients (identifiable)"], r["Primary"], r["Solid PM"],
                   "; ".join(other) or "—", r["Patients with primary + solid PM"]])
table1 = md_table(["Source", "Role", "Specimens", "Patients", "Primary", "Solid PM",
                   "Other material", "Paired patients"], t1rows)

t2raw = read_csv(ROOT / "04_results" / "manuscript" / "Table2_program_effects.csv")
t2rows = []
for r in t2raw:
    ci = f'{num(r["95% CI low"])} to {num(r["95% CI high"])}'
    role = (r["Role"].replace("Alternative definitions of the primary program", "Alternative HRC definition")
            .replace("Control and negative-control programs", "Control")
            .replace("Host-tissue compartments", "Host compartment"))
    t2rows.append([r["Program"], role, num(r["Mean difference"]), ci,
                   pval(r.get("P")), pval(r.get("q"))])
table2 = md_table(["Programme", "Role", "Mean difference", "95% CI", "P", "q"], t2rows,
                  ["---", "---", "---:", "---", "---:", "---:"])

t3raw = read_csv(ROOT / "04_results" / "manuscript" / "Table3_composition_equivalence.csv")
t3rows = []
for r in t3raw:
    ci = f'{num(r["Implied CI low"])} to {num(r["Implied CI high"])}'
    t3rows.append([r["Program"], num(r["Observed difference"], 4),
                   num(r["Slope vs epithelial fraction"]),
                   num(r["Implied epithelial-fraction difference"]), ci])
table3 = md_table(["Programme", "Observed difference", "Slope vs epithelial fraction",
                   "Implied epithelial-fraction difference", "95% CI"], t3rows,
                  ["---", "---:", "---:", "---:", "---"])

man = (SRC / "manuscript_FIG.md").read_text(encoding="utf-8")
man = man.replace("{{TABLE1}}", table1).replace("{{TABLE2}}", table2).replace("{{TABLE3}}", table3)
man_md = OUT / "01_Main_Manuscript_FIG.md"
man_md.write_text(man, encoding="utf-8")

def section(text, start, end):
    return text.split(start, 1)[1].split(end, 1)[0].strip()

abstract = section(man, "## Abstract", "## Introduction")
main_text = section(man, "## Introduction", "## Statements and Declarations")
word_re = re.compile(r"\b[\w]+(?:[-–'][\w]+)*\b", re.UNICODE)
abstract_n = len(word_re.findall(re.sub(r"[*_~]", "", abstract)))
main_n = len(word_re.findall(re.sub(r"[*_~@{}\[\]]", "", main_text)))
if not (150 <= abstract_n <= 250):
    raise RuntimeError(f"Abstract has {abstract_n} words; journal requires 150–250")

title = (SRC / "title_page_FIG.md").read_text(encoding="utf-8")
title = title.replace("{{ABSTRACT_WORD_COUNT}}", str(abstract_n)).replace("{{MAIN_WORD_COUNT}}", str(main_n))
files = {
    "02_Title_Page_FIG.md": title,
    "03_Cover_Letter_FIG.md": (SRC / "cover_letter_FIG.md").read_text(encoding="utf-8"),
    "04_Supplementary_Information_FIG.md": (SRC / "supplementary_FIG.md").read_text(encoding="utf-8"),
    "05_Submission_Metadata_FIG.md": (SRC / "submission_metadata_FIG.md").read_text(encoding="utf-8"),
}
for name, text in files.items():
    (OUT / name).write_text(text, encoding="utf-8")

def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = 2
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def style_docx(path, kind="manuscript"):
    doc = Document(path)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    add_page_number(sec)

    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            if style.name == "Normal":
                style.font.size = Pt(10 if kind in ("manuscript", "supplement") else 11)
    for name, size in [("Title", 16), ("Heading 1", 13), ("Heading 2", 11), ("Heading 3", 10)]:
        if name in doc.styles:
            s = doc.styles[name]
            s.font.name = "Times New Roman"; s.font.size = Pt(size); s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
            s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for p in doc.paragraphs:
        pf = p.paragraph_format
        if p.style.name == "Normal":
            pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            pf.space_after = Pt(4)
        if p._p.xpath(".//w:drawing"):
            if kind != "supplement":
                pf.page_break_before = True
            pf.keep_with_next = False
            p.alignment = 1
        if kind == "supplement" and p.text.startswith("Supplementary Figure S1"):
            pf.page_break_before = True
            pf.keep_with_next = True
        for run in p.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for table in doc.tables:
        table.style = "Table"
        table.autofit = True
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "B7B7B7")
            borders.append(el)
        tblPr.append(borders)
        for i, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
            for cell in row.cells:
                if i == 0:
                    set_cell_shading(cell, "D9EAF7")
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.name = "Times New Roman"; run.font.size = Pt(7.5)
                        if i == 0: run.font.bold = True

    doc.core_properties.title = "Tissue composition in gastric cancer peritoneal metastasis"
    doc.core_properties.subject = "Research Article for Functional & Integrative Genomics"
    doc.save(path)

def run(cmd, cwd=None):
    subprocess.run([str(x) for x in cmd], cwd=cwd, check=True)

def pandoc_doc(md_name, citations=False, kind="manuscript"):
    md = OUT / md_name
    docx = md.with_suffix(".docx")
    cmd = [PANDOC, md, "-o", docx, "--from", "markdown+pipe_tables+subscript",
           "--resource-path", f"{SRC}:{ROOT}"]
    if citations:
        cmd += ["--citeproc", "--bibliography", BIB, "--csl", CSL]
    run(cmd)
    style_docx(docx, kind=kind)
    return docx

docx_files = [
    pandoc_doc("01_Main_Manuscript_FIG.md", citations=True, kind="manuscript"),
    pandoc_doc("02_Title_Page_FIG.md", kind="title"),
    pandoc_doc("03_Cover_Letter_FIG.md", kind="letter"),
    pandoc_doc("04_Supplementary_Information_FIG.md", kind="supplement"),
    pandoc_doc("05_Submission_Metadata_FIG.md", kind="metadata"),
]

# PDF proofs are rendered by the verification step after this builder. LibreOffice is
# intentionally not invoked here because managed environments may require a separate
# headless-application permission for rendering.

# Submission-quality separate figures (600 dpi LZW TIFF), generated from vector PDFs.
figure_pdfs = {
    "Fig1": ROOT / "05_figures" / "main" / "Figure1_study_design.pdf",
    "Fig2": ROOT / "05_figures" / "main" / "Figure2_primary_endpoint.pdf",
    "Fig3": ROOT / "05_figures" / "main" / "Figure3_program_and_compartment_changes.pdf",
    "Fig4": ROOT / "05_figures" / "main" / "Figure4_composition_explains_the_change.pdf",
    "Fig5": ROOT / "05_figures" / "main" / "Figure5_robustness_external_purity.pdf",
    "Fig6": ROOT / "05_figures" / "main" / "Figure6_cell_level_limits.pdf",
    "FigS1": ROOT / "05_figures" / "supplementary" / "Supplementary_Figure_S1_adjustment.pdf",
}
for stem, pdf in figure_pdfs.items():
    run([PDFTOPPM, "-singlefile", "-r", "600", "-tiff", "-tiffcompression", "lzw",
         pdf, FIGOUT / stem])

shutil.copy2(ROOT / "04_results" / "manuscript" / "supplementary_tables_FIG.xlsx",
             OUT / "06_Supplementary_Tables_FIG.xlsx")
shutil.copy2(ROOT / "04_results" / "manuscript" / "main_tables.xlsx",
             OUT / "07_Main_Tables_Source_FIG.xlsx")

print(f"Built FIG submission: abstract={abstract_n} words; main text={main_n} words")
